import streamlit as st
import pandas as pd
import os
import io
import openpyxl
from openpyxl.chart import LineChart, Reference

from utils import sheets_db, data_fetch, returns, auth

st.set_page_config(page_title="Manage - Portfolio Tracker", layout="wide")
auth.require_password()

st.title("🔧 Manage")

PORTFOLIOS = sheets_db.get_portfolios()
sections = list(PORTFOLIOS.values()) + ["Watchlist ETFs", "Backtest history upload", "Manage Portfolios", "Reorder Items", "Export Chart to Excel"]
section = st.sidebar.radio("Section", sections)

POSITION_COLS = {
    "ticker": st.column_config.TextColumn("Ticker", help="6-digit A-share code, e.g. 600519"),
    "name": st.column_config.TextColumn("Name"),
    "asset_type": st.column_config.SelectboxColumn("Type", options=["stock", "etf"]),
    "weight": st.column_config.NumberColumn("Target weight (%)", min_value=0.0, max_value=100.0, step=0.5),
    "purchase_date": st.column_config.DateColumn("Purchase date"),
}

REBALANCE_OPTIONS = {
    "none": "No rebalancing (buy & hold at target weights)",
    "monthly": "Monthly",
    "quarterly": "Quarterly",
    "semiannual": "Every 6 months",
    "annual": "Annually",
}

def _deduplicate_index_data(data):
    """Helper to remove duplicate dates from DataFrames, Series, or dicts of them."""
    if isinstance(data, pd.DataFrame):
        return data[~data.index.duplicated(keep='last')]
    elif isinstance(data, pd.Series):
        return data[~data.index.duplicated(keep='last')]
    elif isinstance(data, dict):
        for k, v in data.items():
            if isinstance(v, (pd.DataFrame, pd.Series)):
                data[k] = v[~v.index.duplicated(keep='last')]
        return data
    return data

def positions_editor(tab_name: str, label: str):
    st.subheader(f"{label} positions")
    current_freq = sheets_db.get_rebalance_frequency(label)
    chosen_freq = st.selectbox("Rebalancing", options=list(REBALANCE_OPTIONS.keys()), format_func=lambda k: REBALANCE_OPTIONS[k], index=list(REBALANCE_OPTIONS.keys()).index(current_freq), key=f"rebal_{tab_name}")
    if chosen_freq != current_freq:
        sheets_db.save_rebalance_frequency(label, chosen_freq)
        st.success(f"Rebalancing set to: {REBALANCE_OPTIONS[chosen_freq]}")
        st.rerun()

    df = sheets_db.read_df(tab_name)
    for col in POSITION_COLS:
        if col not in df.columns:
            df[col] = None
    if not df.empty and "ticker" in df.columns:
        df["ticker"] = df["ticker"].astype(str).str.strip()
    if not df.empty:
        df["purchase_date"] = pd.to_datetime(df["purchase_date"], errors="coerce").dt.date
        df["weight"] = pd.to_numeric(df["weight"], errors="coerce")
    if not df.empty:
        total_weight = df["weight"].sum()
        if abs(total_weight - 100) > 0.5:
            st.warning(f"Weights sum to {total_weight:.1f}% – adjust to 100% for accurate tracking.")
        else:
            st.caption(f"Weights sum to {total_weight:.1f}%. ✅")
    else:
        st.caption("No positions yet – add rows below using the editor.")

    edited = st.data_editor(df[list(POSITION_COLS.keys())], column_config=POSITION_COLS, num_rows="dynamic", use_container_width=True, key=f"editor_{tab_name}")

    if st.button("Save changes", key=f"save_{tab_name}"):
        clean = edited.dropna(subset=["ticker"]).copy()
        clean["ticker"] = clean["ticker"].astype(str).str.strip()
        # FIX: Safely convert dates to strings, replacing NaT with empty string
        clean["purchase_date"] = pd.to_datetime(clean["purchase_date"], errors="coerce").dt.strftime('%Y-%m-%d').fillna('')
        sheets_db.write_df(tab_name, clean)
        sheets_db.clear_caches()
        st.success("Saved.")
        st.rerun()

    if not df.empty:
        st.divider()
        st.subheader("⚖️ Rebalance (save live performance to backtest)")
        
        # FIX: Add a date input so you can specify when the old portfolio ends
        default_rebal_date = pd.Timestamp.now().normalize()
        rebal_date = st.date_input("Rebalance Date", value=default_rebal_date, key=f"rebal_date_{tab_name}")
        rebal_date = pd.Timestamp(rebal_date)
        
        total_weight = df["weight"].sum()
        missing_dates = df["purchase_date"].isna().any()
        weights_valid = abs(total_weight - 100) <= 0.5
        
        if missing_dates:
            st.error("Please ensure all positions have a Purchase Date before rebalancing.")
        elif not weights_valid:
            st.error(f"Weights must sum to 100% to rebalance. Currently sum to {total_weight:.1f}%.")

        if st.button(f"Rebalance {label}", key=f"rebalance_{tab_name}", disabled=(not weights_valid or missing_dates)):
            holdings = []
            for _, row in df.iterrows():
                try:
                    holdings.append({"ticker": str(row["ticker"]).strip(), "asset_type": str(row.get("asset_type", "stock")).strip().lower() or "stock", "weight": float(row["weight"]) / 100.0, "inception_date": pd.to_datetime(row["purchase_date"])})
                except (KeyError, ValueError, TypeError):
                    continue
            if not holdings:
                st.warning("No valid positions to rebalance.")
            else:
                price_data = data_fetch.get_prices_batch(holdings)
                # FIX: Sanitize price data to remove duplicate dates
                price_data = _deduplicate_index_data(price_data)
                
                backtest_index_values = load_backtest(label)
                rebalance_freq = sheets_db.get_rebalance_frequency(label)
                live_start_date = backtest_index_values.index[-1] if not backtest_index_values.empty else None
                
                # FIX: Pass rebal_date as end_date so it only calculates up to that date
                live_index = returns.compute_live_index(
                    holdings, price_data, 
                    rebalance_frequency=rebalance_freq, 
                    live_start_date=live_start_date,
                    end_date=rebal_date
                )
                combined = returns.chain_link_backtest(backtest_index_values, live_index)
                
                # FIX: Drop duplicates AND remove any dates after the Rebalance Date
                if not combined.empty:
                    combined = combined[~combined.index.duplicated(keep='last')]
                    combined = combined[combined.index <= rebal_date]
                    
                if combined.empty:
                    st.warning("Could not compute combined index.")
                else:
                    rebalance_df = combined.reset_index()
                    rebalance_df.columns = ["date", "index_value"]
                    rebalance_df["portfolio"] = label
                    rebalance_df["date"] = pd.to_datetime(rebalance_df["date"], errors="coerce").dt.strftime("%Y-%m-%d").fillna('')
                    rebalance_df["index_value"] = pd.to_numeric(rebalance_df["index_value"], errors="coerce")
                    existing = sheets_db.read_df("backtest_history")
                    if not existing.empty:
                        existing = existing[existing["portfolio"] != label]
                    combined_df = pd.concat([existing, rebalance_df[["date", "portfolio", "index_value"]]], ignore_index=True)
                    sheets_db.write_df("backtest_history", combined_df)
                    sheets_db.clear_caches()
                    st.success(f"✅ Rebalance complete! {label} backtest now includes performance up to {rebal_date.strftime('%Y-%m-%d')}.")
                    st.rerun()

def load_backtest(portfolio_label: str) -> pd.Series:
    df = sheets_db.read_df("backtest_history")
    if df.empty:
        return pd.Series(dtype=float)
    df = df[df["portfolio"] == portfolio_label].copy()
    if df.empty:
        return pd.Series(dtype=float)
    df["date"] = pd.to_datetime(df["date"])
    df = df.sort_values("date")
    # FIX: Drop duplicate dates to prevent reindex errors
    df = df.drop_duplicates(subset=["date"], keep="last")
    return pd.Series(pd.to_numeric(df["index_value"], errors="coerce").values, index=df["date"])

# ----------------------------------------------------------- Excel Export Helpers --
def load_holdings_export(tab_name: str) -> list[dict]:
    df = sheets_db.read_df(tab_name)
    holdings = []
    for _, row in df.iterrows():
        try:
            holdings.append({
                "ticker": str(row["ticker"]).strip(),
                "asset_type": str(row.get("asset_type", "stock")).strip().lower() or "stock",
                "weight": float(row["weight"]) / 100.0,
                "inception_date": pd.to_datetime(row["purchase_date"]),
            })
        except (KeyError, ValueError, TypeError):
            continue
    return holdings

@st.cache_data(ttl=3600, show_spinner=False)
def _compute_portfolio_index_cached(tab_name: str, portfolio_label: str, holdings_key: tuple) -> pd.Series:
    holdings = [
        {"ticker": t, "asset_type": at, "weight": w, "inception_date": pd.Timestamp(d)}
        for t, at, w, d in holdings_key
    ]
    price_data = data_fetch.get_prices_batch(holdings)
    # FIX: Sanitize price data
    price_data = _deduplicate_index_data(price_data)
    
    backtest_index_values = load_backtest(portfolio_label)
    rebalance_freq = sheets_db.get_rebalance_frequency(portfolio_label)
    live_start_date = backtest_index_values.index[-1] if not backtest_index_values.empty else None

    live_index = returns.compute_live_index(
        holdings, price_data,
        rebalance_frequency=rebalance_freq,
        live_start_date=live_start_date,
    )
    if live_index.empty and holdings:
        live_index = returns.compute_live_index(
            holdings, price_data,
            rebalance_frequency=rebalance_freq,
            live_start_date=None,
        )
        if not live_index.empty and live_start_date is not None:
            live_index = live_index[live_index.index >= live_start_date]
            
    combined = returns.chain_link_backtest(backtest_index_values, live_index)
    # FIX: Drop duplicates in final combined output
    if not combined.empty:
        combined = combined[~combined.index.duplicated(keep='last')]
    return combined

def compute_portfolio_index_export(tab_name: str, portfolio_label: str, holdings: list[dict]) -> pd.Series:
    holdings_key = tuple(
        (h["ticker"], h["asset_type"], h["weight"], pd.Timestamp(h["inception_date"]))
        for h in holdings
    )
    return _compute_portfolio_index_cached(tab_name, portfolio_label, holdings_key)

def build_series_options_export() -> dict:
    portfolio_labels = sheets_db.get_portfolios()
    backtest_df = sheets_db.read_df("backtest_history")
    series_options = {}
    for tab_name, label in portfolio_labels.items():
        holdings = load_holdings_export(tab_name)
        if holdings or not backtest_df[backtest_df["portfolio"] == label].empty:
            series_options[label] = compute_portfolio_index_export(tab_name, label, holdings)
    watchlist_df = sheets_db.read_df("watchlist_etfs")
    if not watchlist_df.empty:
        watchlist_prices = data_fetch.get_watchlist_prices(watchlist_df)
        series_options.update(watchlist_prices)
    order_map = sheets_db.get_display_order()
    sorted_keys = sorted(series_options.keys(), key=lambda x: order_map.get(x, 9999))
    return {k: series_options[k] for k in sorted_keys}

def slice_by_period(chart_series: pd.Series, period: str) -> pd.Series:
    chart_series = chart_series.dropna()
    if chart_series.empty:
        return chart_series
    last_date = chart_series.index.max()
    if period == "5D":
        start_date = last_date - pd.Timedelta(days=5)
    elif period == "1M":
        start_date = last_date - pd.DateOffset(months=1)
    elif period == "3M":
        start_date = last_date - pd.DateOffset(months=3)
    elif period == "6M":
        start_date = last_date - pd.DateOffset(months=6)
    elif period == "YTD":
        start_date = pd.Timestamp(year=last_date.year, month=1, day=1)
    elif period == "1Y":
        start_date = last_date - pd.DateOffset(years=1)
    elif period == "2Y":
        start_date = last_date - pd.DateOffset(years=2)
    elif period == "3Y":
        start_date = last_date - pd.DateOffset(years=3)
    elif period == "5Y":
        start_date = last_date - pd.DateOffset(years=5)
    else:
        start_date = chart_series.index.min()
    return chart_series[chart_series.index >= start_date]

def build_excel_bytes(label: str, chart_series: pd.Series, period: str) -> bytes:
    view = slice_by_period(chart_series, period)
    if view.empty or len(view) < 2:
        raise ValueError("Not enough data to export for the selected period.")
    rebased = (view / view.iloc[0] - 1) * 100
    df = pd.DataFrame({
        "Date": view.index,
        "Index Value": view.values,
        "Total Return (%)": rebased.values,
    })
    df["Date"] = pd.to_datetime(df["Date"]).dt.date

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chart Data"
    ws.append(["Date", "Index Value", "Total Return (%)"])
    for _, row in df.iterrows():
        ws.append([row["Date"], float(row["Index Value"]), round(float(row["Total Return (%)"]), 4)])

    for r in range(2, len(df) + 2):
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"

    chart = LineChart()
    chart.title = f"{label} - Total Return % ({period})"
    chart.y_axis.title = "Total return (%)"
    chart.x_axis.title = "Date"
    chart.height = 12
    chart.width = 24

    data_ref = Reference(ws, min_col=3, min_row=1, max_row=len(df) + 1, max_col=3)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=len(df) + 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.legend = None

    ws_chart = wb.create_sheet("Chart")
    ws_chart.add_chart(chart, "A1")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
# ------------------------------------------------------------------------------------

if section in PORTFOLIOS.values():
    tab_name = [k for k, v in PORTFOLIOS.items() if v == section][0]
    positions_editor(tab_name, section)

elif section == "Watchlist ETFs":
    st.subheader("Watchlist ETFs")
    df = sheets_db.read_df("watchlist_etfs")
    cols = {"ticker": st.column_config.TextColumn("Ticker", help="e.g. 510300"), "name": st.column_config.TextColumn("Name")}
    for col in cols:
        if col not in df.columns:
            df[col] = None
    if not df.empty and "ticker" in df.columns:
        df["ticker"] = df["ticker"].astype(str).str.strip()
    edited = st.data_editor(df[list(cols.keys())], column_config=cols, num_rows="dynamic", use_container_width=True, key="editor_watchlist")
    if st.button("Save changes", key="save_watchlist"):
        clean = edited.dropna(subset=["ticker"]).copy()
        clean["ticker"] = clean["ticker"].astype(str).str.strip()
        sheets_db.write_df("watchlist_etfs", clean)
        sheets_db.clear_caches()
        st.success("Saved.")
        st.rerun()

elif section == "Backtest history upload":
    st.subheader("Upload historical backtest returns")
    st.caption("Upload an Excel file with columns: Date, Portfolio, Index Value (starting at 100).")
    allowed_portfolios = list(PORTFOLIOS.values())
    st.write(f"**Allowed Portfolio names:** {', '.join(allowed_portfolios)}")
    os.makedirs("data", exist_ok=True)
    template_path = "data/backtest_template.xlsx"
    if not os.path.exists(template_path):
        pd.DataFrame(columns=["Date", "Portfolio", "Index Value"]).to_excel(template_path, index=False, sheet_name="Backtest Data")
    with open(template_path, "rb") as f:
        st.download_button("Download blank template", f, file_name="backtest_template.xlsx")

    uploaded = st.file_uploader("Upload filled-in template", type=["xlsx"])
    if uploaded is not None:
        try:
            new_data = pd.read_excel(uploaded, sheet_name="Backtest Data", dtype=str)
            if "Index Value" in new_data.columns:
                new_data["Index Value"] = new_data["Index Value"].str.replace(",", ".").str.replace(" ", "").str.replace("'", "")
                new_data["Index Value"] = new_data["Index Value"].str.replace(r"[^\d.\-]", "", regex=True)
                new_data["Index Value"] = pd.to_numeric(new_data["Index Value"], errors="coerce")
            if "Date" in new_data.columns:
                new_data["Date"] = pd.to_datetime(new_data["Date"], errors="coerce")
            new_data = new_data.dropna(subset=["Date", "Index Value", "Portfolio"])
        except Exception as e:
            st.error(f"Couldn't read that file: {e}")
            new_data = None

        if new_data is not None:
            required = {"Date", "Portfolio", "Index Value"}
            if not required.issubset(new_data.columns):
                st.error(f"Missing columns. Found: {list(new_data.columns)}")
            else:
                bad = set(new_data["Portfolio"]) - set(allowed_portfolios)
                if bad:
                    st.error(f"Unrecognized portfolio(s): {bad}. Must match an existing portfolio name exactly.")
                else:
                    st.dataframe(new_data, use_container_width=True)
                    if st.button("Confirm and save"):
                        existing = sheets_db.read_df("backtest_history")
                        uploaded_pf = set(new_data["Portfolio"])
                        if not existing.empty:
                            existing = existing[~existing["portfolio"].isin(uploaded_pf)]
                        new_data = new_data.rename(columns={"Date": "date", "Portfolio": "portfolio", "Index Value": "index_value"})
                        new_data["date"] = pd.to_datetime(new_data["date"], errors="coerce").dt.strftime("%Y-%m-%d").fillna('')
                        new_data["index_value"] = pd.to_numeric(new_data["index_value"], errors="coerce")
                        combined_df = pd.concat([existing, new_data[["date", "portfolio", "index_value"]]], ignore_index=True)
                        sheets_db.write_df("backtest_history", combined_df)
                        sheets_db.clear_caches()
                        st.success("Backtest history saved.")
                        st.rerun()

    st.divider()
    st.write("Current stored backtest history:")
    st.dataframe(sheets_db.read_df("backtest_history"), use_container_width=True)

elif section == "Manage Portfolios":
    st.subheader("Manage Portfolios")
    st.caption("Add new portfolios or delete existing ones. Deleted portfolios cannot be recovered.")
    with st.form("add_portfolio_form"):
        new_label = st.text_input("New Portfolio Name")
        submitted = st.form_submit_button("Add Portfolio")
        if submitted and new_label:
            sheets_db.add_portfolio(new_label)
            st.success(f"Added portfolio: {new_label}")
            st.rerun()
    st.divider()
    st.write("**Existing Portfolios:**")
    for tab, label in PORTFOLIOS.items():
        col1, col2 = st.columns([4, 1])
        col1.write(f"{label} (`{tab}`)")
        if col2.button("Delete", key=f"del_{tab}"):
            sheets_db.delete_portfolio(tab, label)
            st.warning(f"Deleted {label}")
            st.rerun()

elif section == "Reorder Items":
    st.subheader("Reorder Portfolios & ETFs")
    st.caption("Set the display order for the main dashboard dropdown and comparison table.")
    all_items = list(PORTFOLIOS.values())
    watchlist = sheets_db.read_df("watchlist_etfs")
    if not watchlist.empty:
        for _, row in watchlist.iterrows():
            name = str(row.get("name", "")).strip() or row["ticker"]
            all_items.append(f"{name} ({row['ticker']})")
    if not all_items:
        st.info("No portfolios or ETFs found to reorder.")
    else:
        current_order = sheets_db.get_display_order()
        order_list = [current_order.get(item, 99) for item in all_items]
        df_order = pd.DataFrame({"Item": all_items, "Sort Order": order_list})
        edited = st.data_editor(df_order, num_rows="fixed", use_container_width=True, key="order_editor")
        if st.button("Save Order"):
            sorted_df = edited.sort_values("Sort Order")
            sheets_db.save_display_order(sorted_df["Item"].tolist())
            st.success("Display order saved!")
            st.rerun()

elif section == "Export Chart to Excel":
    st.subheader("📊 Export chart to Excel")
    st.caption(
        "Pick a portfolio or ETF and a period. The Excel file will contain "
        "the underlying data series on one sheet and a native Excel line chart "
        "on another sheet."
    )

    try:
        series_options = build_series_options_export()
    except Exception as e:
        st.error(f"Failed to load series: {e}")
        series_options = {}

    if not series_options:
        st.info("No portfolios or watchlist ETFs set up yet. Add some first.")
    else:
        col1, col2 = st.columns([3, 2])
        with col1:
            choice = st.selectbox("Choose what to chart:", list(series_options.keys()), key="export_choice")
        with col2:
            period = st.selectbox(
                "Period",
                options=["5D", "1M", "3M", "6M", "YTD", "1Y", "2Y", "3Y", "5Y", "Max"],
                index=9,
                key="export_period",
            )

        chart_series = series_options.get(choice)
        if chart_series is None or chart_series.dropna().empty:
            st.warning("No data available for this selection.")
        else:
            try:
                excel_bytes = build_excel_bytes(choice, chart_series, period)
                safe_name = "".join(c if c.isalnum() or c in "-_" else "_" for c in choice)
                file_name = f"{safe_name}_{period}.xlsx"
                st.download_button(
                    label="⬇️ Download Excel file",
                    data=excel_bytes,
                    file_name=file_name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                st.success("Excel file is ready — click the button above to download.")
            except Exception as e:
                st.error(f"Could not build Excel: {e}")
